import tkinter as tk
from tkinter import ttk, filedialog, messagebox
import pandas as pd
from openpyxl import load_workbook
from openpyxl.utils.dataframe import dataframe_to_rows

class SheetSegregator(tk.Tk):
    def __init__(self):
        super().__init__()
        self.title("Excel Sheet Segregator")
        self.resizable(False, False)
        self.filename = None
        self.excel = None
        self.df = None

        # --- File picker ---
        file_frame = ttk.Frame(self, padding=10)
        file_frame.grid(row=0, column=0, sticky="ew")
        ttk.Label(file_frame, text="1) Choose workbook:").grid(row=0, column=0, sticky="w")
        self.file_lbl = ttk.Label(file_frame, text="No file selected", width=40)
        self.file_lbl.grid(row=0, column=1, padx=5)
        ttk.Button(file_frame, text="Browse…", command=self.load_file).grid(row=0, column=2)

        # --- Sheet selector ---
        sheet_frame = ttk.Frame(self, padding=10)
        sheet_frame.grid(row=1, column=0, sticky="ew")
        ttk.Label(sheet_frame, text="2) Select sheet:").grid(row=0, column=0, sticky="w")
        self.sheet_cb = ttk.Combobox(sheet_frame, state="disabled", width=37,
                                     postcommand=self.update_sheets)
        self.sheet_cb.grid(row=0, column=1, padx=5)
        self.sheet_cb.bind("<<ComboboxSelected>>", lambda e: self.load_df())

        # --- Column selector ---
        col_frame = ttk.Frame(self, padding=10)
        col_frame.grid(row=2, column=0, sticky="ew")
        ttk.Label(col_frame, text="3) Select column:").grid(row=0, column=0, sticky="w")
        self.col_cb = ttk.Combobox(col_frame, state="disabled", width=37)
        self.col_cb.grid(row=0, column=1, padx=5)

        # --- Segregate button ---
        btn_frame = ttk.Frame(self, padding=10)
        btn_frame.grid(row=3, column=0, sticky="ew")
        self.go_btn = ttk.Button(btn_frame, text="Segregate →", state="disabled",
                                 command=self.segregate)
        self.go_btn.grid(row=0, column=1, sticky="e")

    def load_file(self):
        fn = filedialog.askopenfilename(
            filetypes=[("Excel files","*.xlsx;*.xlsm;*.xls")])
        if not fn:
            return
        self.filename = fn
        self.file_lbl.config(text=fn.split("/")[-1])
        self.sheet_cb.config(state="readonly")
        try:
            self.excel = pd.ExcelFile(self.filename, engine="openpyxl")
            self.sheet_cb['values'] = self.excel.sheet_names
        except Exception as e:
            messagebox.showerror("Error", f"Could not open file:\n{e}")
            self.sheet_cb.config(state="disabled")

    def update_sheets(self):
        if self.filename:
            try:
                self.excel = pd.ExcelFile(self.filename, engine="openpyxl")
                self.sheet_cb['values'] = self.excel.sheet_names
            except:
                pass

    def load_df(self):
        sheet = self.sheet_cb.get()
        if not sheet:
            return
        try:
            self.df = pd.read_excel(self.filename, sheet_name=sheet, engine="openpyxl")
        except Exception as e:
            messagebox.showerror("Error", f"Could not read sheet:\n{e}")
            return
        cols = list(self.df.columns)
        self.col_cb.config(state="readonly", values=cols)
        self.col_cb.set('')
        self.go_btn.config(state="disabled")
        self.col_cb.bind("<<ComboboxSelected>>", lambda e: self.go_btn.config(state="normal"))
    def segregate(self):
        col = self.col_cb.get()
        if not col:
            messagebox.showwarning("Select Column", "Please select a column first.")
            return

        # index of the column to filter on
        idx = list(self.df.columns).index(col)

        unique_vals = self.df[col].dropna().unique()
        if len(unique_vals) == 0:
            messagebox.showinfo("Nothing to do", f"No values found in '{col}'.")
            return

        try:
            wb = load_workbook(self.filename)
            orig_ws = wb[self.sheet_cb.get()]

            # copy column widths
            col_widths = {
                col_letter: dim.width
                for col_letter, dim in orig_ws.column_dimensions.items()
                if dim.width is not None
            }
            # copy row heights
            row_heights = {
                row_idx: dim.height
                for row_idx, dim in orig_ws.row_dimensions.items()
                if dim.height is not None
            }

            for val in unique_vals:
                name = str(val)[:31]
                # remove old sheet
                if name in wb.sheetnames:
                    wb.remove(wb[name])
                new_ws = wb.create_sheet(name)

                # apply col widths to new sheet
                for col_letter, w in col_widths.items():
                    new_ws.column_dimensions[col_letter].width = w
                # apply row heights (we'll reapply per-row as we go)

                # --- copy header row (row 1) ---
                for src_cell in orig_ws[1]:
                    dest = new_ws.cell(row=1, column=src_cell.column)
                    # copy value & style
                    dest.value = src_cell.value
                    dest.font = src_cell.font.copy()
                    dest.fill = src_cell.fill.copy()
                    dest.border = src_cell.border.copy()
                    dest.alignment = src_cell.alignment.copy()
                    dest.number_format = src_cell.number_format
                # header row height
                if 1 in row_heights:
                    new_ws.row_dimensions[1].height = row_heights[1]

                # --- copy data rows matching this value ---
                dest_row = 2
                for src_row in orig_ws.iter_rows(min_row=2, values_only=False):
                    if src_row[idx].value == val:
                        # copy each cell in the row
                        for src_cell in src_row:
                            dest = new_ws.cell(row=dest_row, column=src_cell.column)
                            dest.value = src_cell.value
                            dest.font = src_cell.font.copy()
                            dest.fill = src_cell.fill.copy()
                            dest.border = src_cell.border.copy()
                            dest.alignment = src_cell.alignment.copy()
                            dest.number_format = src_cell.number_format
                        # apply that row’s height
                        if src_cell.row in row_heights:
                            new_ws.row_dimensions[dest_row].height = row_heights[src_cell.row]
                        dest_row += 1

            wb.save(self.filename)
            wb.close()
            messagebox.showinfo("Done ✅",
                                f"Created {len(unique_vals)} sheet(s) for column '{col}',\n"
                                "with formatting, widths, heights, and borders preserved.")
        except Exception as e:
            messagebox.showerror("Error", f"Failed to write sheets:\n{e}")

if __name__ == "__main__":
    app = SheetSegregator()
    app.mainloop()
